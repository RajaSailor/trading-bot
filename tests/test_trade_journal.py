import json
import importlib
import sqlite3
import zipfile
from datetime import datetime
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from tools.sync_trade_journal import sync_workbook
from trade_journal import (
    HEADER_ROW,
    ROW_ID_COLUMN,
    SHEET_NAMES,
    TradeJournalExporter,
    TradeJournalWorker,
)

IST = ZoneInfo("Asia/Kolkata")
TOKEN = "test-only-journal-export-token-0123456789"


def _epoch(hour, minute, second=0):
    return datetime(2026, 10, 14, hour, minute, second, tzinfo=IST).timestamp()


def _signal_db(path, positions):
    with sqlite3.connect(path) as db:
        db.execute("""
            CREATE TABLE lifecycle_signals(
                id TEXT PRIMARY KEY, category TEXT, status TEXT, opened_at REAL,
                closed_at REAL, last_quote_at REAL, data TEXT
            )
        """)
        for item in positions:
            db.execute(
                "INSERT INTO lifecycle_signals VALUES(?,?,?,?,?,?,?)",
                (
                    item["id"], item["category"], item["status"], item["opened_at"],
                    item.get("closed_at"), item.get("closed_at") or item["opened_at"],
                    json.dumps(item),
                ),
            )


def _live_db(path, items):
    with sqlite3.connect(path) as db:
        db.execute(
            "CREATE TABLE IF NOT EXISTS live_intents(id TEXT PRIMARY KEY, signal_key TEXT, data TEXT)"
        )
        db.execute("DELETE FROM live_intents")
        for item in items:
            db.execute(
                "INSERT INTO live_intents VALUES(?,?,?)",
                (item["id"], item.get("signal_key"), json.dumps(item)),
            )


def _position(identifier="sig-1", *, status="OPEN", multiplier=None, contract="NIFTY-CE"):
    opened = _epoch(10, 0)
    position = {
        "id": identifier,
        "category": "index_options",
        "underlying": "NIFTY",
        "status": status,
        "opened_at": opened,
        "closed_at": _epoch(10, 1, 30) if status != "OPEN" else None,
        "entry": 100.0,
        "observed_entry": 100.05,
        "exit_premium": 120.0 if status != "OPEN" else None,
        "strike": 25000,
        "expiry": "2026-10-29",
        "option_type": "CE",
        "option_symbol": contract,
        "lot_size": 50,
        "metadata": {},
    }
    if multiplier is not None:
        position["metadata"]["premium_multiplier"] = multiplier
    return position


def _header_map(sheet):
    return {cell.value: cell.column for cell in sheet[HEADER_ROW]}


def _data_rows(sheet):
    return [
        row for row in range(HEADER_ROW + 1, sheet.max_row + 1)
        if sheet.cell(row, ROW_ID_COLUMN).value
    ]


def _workbook_bytes(path):
    return path.read_bytes()


class _Response:
    def __init__(self, payload, status=200):
        self.payload = payload
        self.status_code = status

    def iter_content(self, chunk_size):
        for offset in range(0, len(self.payload), chunk_size):
            yield self.payload[offset:offset + chunk_size]

    def close(self):
        pass


class _Session:
    def __init__(self, payload, status=200):
        self.payload, self.status = payload, status
        self.headers = None

    def get(self, url, *, headers, timeout, stream, allow_redirects):
        assert allow_redirects is False
        self.headers = headers
        return _Response(self.payload, self.status)


def test_four_sheets_open_to_closed_same_row_and_signal_pnl(tmp_path):
    signals = tmp_path / "signals.db"
    live = tmp_path / "live.db"
    output = tmp_path / "journal.xlsx"
    position = _position(multiplier=1)
    _signal_db(signals, [position])
    exporter = TradeJournalExporter(signals, live, output, clock=lambda: _epoch(12, 0))

    assert exporter.export()
    workbook = load_workbook(output, data_only=False)
    assert tuple(workbook.sheetnames) == SHEET_NAMES
    sheet = workbook["Index"]
    row = _data_rows(sheet)[0]
    columns = _header_map(sheet)
    stable_id = sheet.cell(row, columns["Row ID"]).value
    assert sheet.cell(row, columns["Status"]).value == "OPEN"
    assert sheet.cell(row, columns["Exit Time (IST)"]).value is None
    assert sheet.cell(row, columns["Gross P&L (INR)"]).value is None
    assert sheet.cell(row, columns["Entry Order Type"]).value.startswith("N/A")
    workbook.close()

    position.update(status="TARGET", closed_at=_epoch(10, 1, 30), exit_premium=120)
    with sqlite3.connect(signals) as db:
        db.execute(
            "UPDATE lifecycle_signals SET status=?,closed_at=?,data=? WHERE id=?",
            ("TARGET", position["closed_at"], json.dumps(position), position["id"]),
        )
    assert exporter.export()
    workbook = load_workbook(output, data_only=False)
    sheet = workbook["Index"]
    row = _data_rows(sheet)[0]
    columns = _header_map(sheet)
    assert sheet.cell(row, columns["Row ID"]).value == stable_id
    assert sheet.cell(row, columns["Exit Price"]).value == 120
    assert sheet.cell(row, columns["Total Duration"]).value.total_seconds() == 90
    assert sheet.cell(row, columns["Gross P&L (INR)"]).value == 1000
    assert sheet.cell(row, columns["Target Hit"]).value == "Yes"
    assert sheet.cell(row, columns["Win % (running gross)"]).value == 1
    assert sheet.cell(row, columns["Time Zone"]).value == "IST"
    workbook.close()
    with zipfile.ZipFile(output) as archive:
        assert archive.testzip() is None
    restarted = TradeJournalExporter(signals, live, output)
    restarted.export()
    workbook = load_workbook(output)
    assert len(_data_rows(workbook["Index"])) == 1
    assert workbook["Index"].cell(
        _data_rows(workbook["Index"])[0], columns["Row ID"]
    ).value == stable_id
    workbook.close()


@pytest.mark.parametrize(
    ("status", "target", "stop", "session"),
    [("TARGET", "Yes", "No", "No"), ("STOP", "No", "Yes", "No"),
     ("SESSION_CLOSE", "No", "No", "Yes")],
)
def test_signal_exit_flags_and_never_claims_an_order(status, target, stop, session, tmp_path):
    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    _signal_db(signals, [_position(status=status)])
    exporter = TradeJournalExporter(signals, live, output)
    exporter.export()
    workbook = load_workbook(output)
    sheet = workbook["Index"]
    row, columns = _data_rows(sheet)[0], _header_map(sheet)
    assert sheet.cell(row, columns["Target Hit"]).value == target
    assert sheet.cell(row, columns["Stop Hit"]).value == stop
    assert sheet.cell(row, columns["Session Auto-Square-Off"]).value == session
    assert sheet.cell(row, columns["Entry Order Type"]).value.startswith("N/A")
    assert sheet.cell(row, columns["Exit Order Type"]).value.startswith("N/A")
    workbook.close()


def test_missing_conversion_and_formula_injection_are_explicit(tmp_path):
    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    position = _position(status="TARGET", contract='=HYPERLINK("https://bad.invalid")')
    _signal_db(signals, [position])
    exporter = TradeJournalExporter(signals, live, output)
    exporter.export()
    workbook = load_workbook(output, data_only=False)
    sheet = workbook["Index"]
    row, columns = _data_rows(sheet)[0], _header_map(sheet)
    assert sheet.cell(row, columns["Gross P&L (INR)"]).value is None
    assert "premium_multiplier missing" in sheet.cell(row, columns["Accounting Basis"]).value
    contract = sheet.cell(row, columns["Contract Symbol"])
    assert contract.value.startswith("'=")
    assert contract.data_type != "f"
    assert sheet.cell(row, columns["Net P&L (INR)"]).value is None
    workbook.close()


def test_missing_price_or_lot_never_creates_amounts_or_fake_units(tmp_path):
    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    position = _position(status="TARGET")
    position.update(lot_size=None, entry=None, exit_premium=None)
    _signal_db(signals, [position])
    TradeJournalExporter(signals, live, output).export()
    workbook = load_workbook(output)
    sheet = workbook["Index"]
    row, columns = _data_rows(sheet)[0], _header_map(sheet)
    for column in (
        "Entry Price", "Exit Price", "Gross P&L (INR)", "Net P&L (INR)",
        "Qty Lots", "Lot Size", "Units", "Amount Used (INR)", "Amount After Exit (INR)",
    ):
        assert sheet.cell(row, columns[column]).value is None
    workbook.close()


def test_live_partial_fills_update_same_row_only_when_fully_closed(tmp_path):
    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    _signal_db(signals, [_position()])
    entry = {
        "id": "live-1", "signal_key": "stable-signal", "side": "BUY",
        "entry_notified": True, "filled_quantity": 50, "quantity": 50,
        "average_price": 10, "first_fill_timestamp": _epoch(10, 2),
        "owned_quantity": 25, "lot_size": 50, "lots": 1, "underlying": "NIFTY",
        "strike": 25000, "expiry": "2026-10-29", "option_type": "CE",
        "exit_filled_quantity": 25, "exit_average_price": 12,
        "exit_fill_timestamp": _epoch(10, 3),
    }
    _live_db(live, [entry])
    exporter = TradeJournalExporter(signals, live, output)
    exporter.export()
    workbook = load_workbook(output)
    sheet = workbook["Live Trades"]
    row, columns = _data_rows(sheet)[0], _header_map(sheet)
    row_id = sheet.cell(row, columns["Row ID"]).value
    assert sheet.cell(row, columns["Status"]).value == "OPEN"
    assert sheet.cell(row, columns["Exit Time (IST)"]).value is None
    assert sheet.cell(row, columns["Gross P&L (INR)"]).value is None
    workbook.close()

    entry.update(
        owned_quantity=0,
        exit_filled_quantity=50,
        exit_average_price=12,
        exit_fill_timestamp=_epoch(10, 4),
        exit_reason="time_exit_1525",
        time_exit=True,
        legs_triggered=["TARGET_LEG"],
    )
    _live_db(live, [entry])
    exporter.export()
    workbook = load_workbook(output)
    sheet = workbook["Live Trades"]
    row, columns = _data_rows(sheet)[0], _header_map(sheet)
    assert sheet.cell(row, columns["Row ID"]).value == row_id
    assert sheet.cell(row, columns["Status"]).value == "CLOSED"
    assert sheet.cell(row, columns["Units"]).value == 50
    assert sheet.cell(row, columns["Gross P&L (INR)"]).value == 100
    assert sheet.cell(row, columns["Net P&L (INR)"]).value is None
    assert sheet.cell(row, columns["Session Auto-Square-Off"]).value == "Yes"
    assert sheet.cell(row, columns["Target Hit"]).value == "No"
    assert sheet.cell(row, columns["Stop Hit"]).value == "No"
    assert sheet.cell(row, columns["Exit Order Type"]).value == "LIMIT"
    assert sheet.cell(row, columns["Charges Provenance/Status"]).value.startswith("Pending")
    workbook.close()


def test_unconfirmed_live_intents_are_not_exported_and_sheets_are_separate(tmp_path):
    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    _signal_db(signals, [_position()])
    _live_db(live, [{
        "id": "not-filled", "side": "BUY", "filled_quantity": 50,
        "entry_notified": False, "average_price": 10,
    }])
    TradeJournalExporter(signals, live, output).export()
    workbook = load_workbook(output)
    assert len(_data_rows(workbook["Index"])) == 1
    assert len(_data_rows(workbook["Live Trades"])) == 0
    workbook.close()


def test_manual_live_exit_uses_weighted_broker_fills_once(tmp_path):
    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    _signal_db(signals, [])
    parent = {
        "id": "live-manual", "signal_key": "stable-signal", "side": "BUY",
        "entry_notified": True, "filled_quantity": 50, "quantity": 50,
        "average_price": 10.5, "first_fill_timestamp": _epoch(10, 2),
        "owned_quantity": 0, "lot_size": 25, "lots": 2, "underlying": "NIFTY",
        "strike": 25000, "expiry": "2026-10-29", "option_type": "PE",
    }
    exits = [
        {
            "id": "sell-1", "side": "SELL", "entry_id": "live-manual",
            "filled_quantity": 20, "average_price": 12, "fill_timestamp": _epoch(10, 3),
            "order_type": "LIMIT",
        },
        {
            "id": "sell-2", "side": "SELL", "entry_id": "live-manual",
            "filled_quantity": 30, "average_price": 14, "fill_timestamp": _epoch(10, 4),
            "order_type": "LIMIT",
        },
    ]
    _live_db(live, [parent, *exits])
    TradeJournalExporter(signals, live, output).export()
    workbook = load_workbook(output)
    sheet = workbook["Live Trades"]
    rows, columns = _data_rows(sheet), _header_map(sheet)
    assert len(rows) == 1
    row = rows[0]
    assert sheet.cell(row, columns["Exit Price"]).value == pytest.approx(13.2)
    assert sheet.cell(row, columns["Gross P&L (INR)"]).value == pytest.approx(135)
    assert sheet.cell(row, columns["Units"]).value == 50
    assert sheet.cell(row, columns["Qty Lots"]).value == 2
    assert sheet.cell(row, columns["Exit Order Type"]).value == "LIMIT"
    assert sheet.cell(row, columns["Target Hit"]).value == "No"
    workbook.close()


def test_daily_win_rate_counts_breakeven_and_excludes_open_unpriced(tmp_path):
    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    positions = [
        _position("win", status="TARGET", multiplier=1),
        _position("loss", status="STOP", multiplier=1),
        _position("flat", status="SESSION_CLOSE", multiplier=1),
        _position("unpriced", status="TARGET"),
        _position("open", status="OPEN", multiplier=1),
    ]
    positions[1]["exit_premium"] = 90
    positions[2]["exit_premium"] = 100
    positions[3]["exit_premium"] = 120
    positions[4]["closed_at"] = None
    stock = _position("stock", status="TARGET", multiplier=1)
    stock["category"] = "nifty50_stock_options"
    commodity = _position("commodity", status="SESSION_CLOSE")
    commodity["category"] = "commodity_options"
    positions.extend((stock, commodity))
    _signal_db(signals, positions)
    TradeJournalExporter(signals, live, output).export()
    workbook = load_workbook(output)
    index = workbook["Index"]
    summary = index.cell(4, 1).value
    assert "wins 1; losses 1; break-even 1; unpriced 1" in summary
    assert "gross win rate 33.3%" in summary
    assert len(_data_rows(index)) == 5
    assert len(_data_rows(workbook["Stocks"])) == 1
    assert len(_data_rows(workbook["Commodity"])) == 1
    assert len(_data_rows(workbook["Live Trades"])) == 0
    workbook.close()


def test_local_sync_auth_failure_and_xls_refusal_leave_old_file(tmp_path):
    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    _signal_db(signals, [_position(status="TARGET", multiplier=1)])
    TradeJournalExporter(signals, live, output).export()
    previous = _workbook_bytes(output)
    failed = _Session(b"unauthorized", status=401)
    with pytest.raises(RuntimeError):
        sync_workbook("https://bot.example/trade-journal.xlsx", TOKEN, output, session=failed)
    assert _workbook_bytes(output) == previous
    assert failed.headers == {"X-Trade-Journal-Token": TOKEN}
    with pytest.raises(ValueError, match="\\.xls"):
        sync_workbook("https://bot.example/trade-journal.xlsx", TOKEN, tmp_path / "journal.xls")
    with pytest.raises(ValueError, match="query strings"):
        sync_workbook("https://bot.example/trade-journal.xlsx?token=x", TOKEN, output)


def test_local_sync_retries_transient_file_lock_without_losing_old_workbook(tmp_path, monkeypatch):
    import tools.sync_trade_journal as sync_module

    signals, live = tmp_path / "signals.db", tmp_path / "live.db"
    output, source = tmp_path / "journal.xlsx", tmp_path / "source.xlsx"
    _signal_db(signals, [_position(status="TARGET", multiplier=1)])
    TradeJournalExporter(signals, live, source).export()
    output.write_bytes(_workbook_bytes(source))
    original = sync_module.os.replace
    attempts = []

    def replace(source_path, target_path):
        attempts.append((source_path, target_path))
        if len(attempts) < 3:
            raise PermissionError("simulated Excel lock")
        return original(source_path, target_path)

    monkeypatch.setattr(sync_module.os, "replace", replace)
    sync_workbook(
        "https://bot.example/trade-journal.xlsx", TOKEN, output,
        session=_Session(_workbook_bytes(source)), sleep=lambda _seconds: None,
    )
    assert len(attempts) == 3
    with zipfile.ZipFile(output) as archive:
        assert archive.testzip() is None


def test_sync_preserves_review_fields_and_creates_backup(tmp_path):
    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    _signal_db(signals, [_position(status="TARGET", multiplier=1)])
    exporter = TradeJournalExporter(signals, live, tmp_path / "source.xlsx")
    exporter.export()
    first_download = _Session(_workbook_bytes(tmp_path / "source.xlsx"))
    sync_workbook("https://bot.example/trade-journal.xlsx", TOKEN, output, session=first_download)

    workbook = load_workbook(output)
    sheet = workbook["Index"]
    row, columns = _data_rows(sheet)[0], _header_map(sheet)
    sheet.cell(row, columns["Notes"]).value = "reviewed"
    sheet.cell(row, columns["Review"]).value = "follow up"
    workbook.save(output)
    workbook.close()

    exporter.export()
    sync_workbook(
        "https://bot.example/trade-journal.xlsx", TOKEN, output,
        session=_Session(_workbook_bytes(tmp_path / "source.xlsx")),
    )
    workbook = load_workbook(output)
    sheet = workbook["Index"]
    row, columns = _data_rows(sheet)[0], _header_map(sheet)
    assert sheet.cell(row, columns["Notes"]).value == "reviewed"
    assert sheet.cell(row, columns["Review"]).value == "follow up"
    workbook.close()
    assert list(tmp_path.glob("journal.*.bak.xlsx"))


def test_sync_refuses_to_drop_user_review_data(tmp_path):
    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    position = _position(status="TARGET", multiplier=1)
    _signal_db(signals, [position])
    source = tmp_path / "source.xlsx"
    exporter = TradeJournalExporter(signals, live, source)
    exporter.export()
    sync_workbook(
        "https://bot.example/trade-journal.xlsx", TOKEN, output,
        session=_Session(_workbook_bytes(source)),
    )
    workbook = load_workbook(output)
    sheet = workbook["Index"]
    row, columns = _data_rows(sheet)[0], _header_map(sheet)
    sheet.cell(row, columns["Notes"]).value = "keep this note"
    workbook.save(output)
    workbook.close()
    previous = _workbook_bytes(output)

    with sqlite3.connect(signals) as db:
        db.execute("DELETE FROM lifecycle_signals WHERE id=?", (position["id"],))
    exporter.export()
    with pytest.raises(ValueError, match="discard user review"):
        sync_workbook(
            "https://bot.example/trade-journal.xlsx", TOKEN, output,
            session=_Session(_workbook_bytes(source)),
        )
    assert _workbook_bytes(output) == previous


def test_sync_rejects_edited_bot_owned_cells_without_replacing_file(tmp_path):
    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    _signal_db(signals, [_position(status="TARGET", multiplier=1)])
    source = tmp_path / "source.xlsx"
    TradeJournalExporter(signals, live, source).export()
    sync_workbook(
        "https://bot.example/trade-journal.xlsx", TOKEN, output,
        session=_Session(_workbook_bytes(source)),
    )
    workbook = load_workbook(output)
    sheet = workbook["Index"]
    row, columns = _data_rows(sheet)[0], _header_map(sheet)
    sheet.cell(row, columns["Entry Price"]).value = 999
    workbook.save(output)
    workbook.close()
    previous = _workbook_bytes(output)
    with pytest.raises(ValueError, match="bot-owned"):
        sync_workbook(
            "https://bot.example/trade-journal.xlsx", TOKEN, output,
            session=_Session(_workbook_bytes(source)),
        )
    assert _workbook_bytes(output) == previous


def test_worker_replays_database_updates_and_preserves_pending_work(tmp_path, monkeypatch):
    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    position = _position()
    _signal_db(signals, [position])
    exporter = TradeJournalExporter(signals, live, output)
    worker = TradeJournalWorker(exporter)
    assert worker.run_once(force=True)
    assert not worker.run_once()

    position.update(status="TARGET", closed_at=_epoch(10, 1), exit_premium=120)
    with sqlite3.connect(signals) as db:
        db.execute(
            "UPDATE lifecycle_signals SET status=?,closed_at=?,data=? WHERE id=?",
            ("TARGET", position["closed_at"], json.dumps(position), position["id"]),
        )
    original = exporter.export
    monkeypatch.setattr(exporter, "export", lambda: False)
    assert not worker.run_once()
    monkeypatch.setattr(exporter, "export", original)
    assert worker.run_once()
    workbook = load_workbook(output)
    sheet = workbook["Index"]
    row, columns = _data_rows(sheet)[0], _header_map(sheet)
    assert sheet.cell(row, columns["Status"]).value == "TARGET"
    workbook.close()


def test_failed_export_does_not_replace_last_good_workbook(tmp_path):
    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    _signal_db(signals, [_position()])
    TradeJournalExporter(signals, live, output).export()
    previous = _workbook_bytes(output)
    broken = TradeJournalExporter(tmp_path / "missing.db", live, output)
    with pytest.raises(FileNotFoundError):
        broken.export()
    assert _workbook_bytes(output) == previous


def test_concurrent_exports_coalesce_and_leave_valid_workbook(tmp_path, monkeypatch):
    from threading import Event, Thread

    signals, live, output = (tmp_path / "signals.db", tmp_path / "live.db", tmp_path / "journal.xlsx")
    _signal_db(signals, [_position()])
    exporter = TradeJournalExporter(signals, live, output)
    entered, release = Event(), Event()
    snapshot = exporter._snapshot

    def blocking_snapshot():
        entered.set()
        assert release.wait(3)
        return snapshot()

    monkeypatch.setattr(exporter, "_snapshot", blocking_snapshot)
    outcomes = []
    thread = Thread(target=lambda: outcomes.append(exporter.export()))
    thread.start()
    assert entered.wait(3)
    assert exporter.export() is False
    release.set()
    thread.join(timeout=5)
    assert not thread.is_alive()
    assert outcomes == [True]
    with zipfile.ZipFile(output) as archive:
        assert archive.testzip() is None


def test_server_download_requires_header_token_and_disables_caching(tmp_path, monkeypatch):
    source = tmp_path / "signals.db"
    live = tmp_path / "live.db"
    output = tmp_path / "journal.xlsx"
    _signal_db(source, [])
    TradeJournalExporter(source, live, output).export()
    monkeypatch.setenv("FLASK_ENV", "development")
    monkeypatch.setenv("TRADE_JOURNAL_EXPORT_TOKEN", TOKEN)
    import main
    module = importlib.reload(main)
    monkeypatch.setattr(
        module, "trade_journal_worker",
        SimpleNamespace(exporter=SimpleNamespace(output_path=output)),
    )
    client = module.app.test_client()
    assert client.get("/trade-journal.xlsx").status_code == 401
    response = client.get(
        "/trade-journal.xlsx", headers={"X-Trade-Journal-Token": TOKEN}
    )
    assert response.status_code == 200
    assert response.headers["Cache-Control"] == "private, no-store"
    assert response.headers["X-Content-Type-Options"] == "nosniff"
    assert response.data.startswith(b"PK")
