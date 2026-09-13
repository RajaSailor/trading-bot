from __future__ import annotations

import threading
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook, load_workbook


@dataclass
class TradeRequest:
    trade_id: str
    symbol: str
    signal: str
    entry: float
    stop_loss: float
    targets: List[float]
    category: str
    option_symbol: str
    option_type: str
    quantity: int = 1
    order_type: str = "STOPLOSS_LIMIT_BUY"
    validity: str = "INTRADAY"
    mode: str = "REGULAR"
    status: str = "PENDING"
    created_at: str = field(default_factory=lambda: datetime.now(UTC).isoformat())
    approved_at: Optional[str] = None
    executed_at: Optional[str] = None
    rejection_reason: Optional[str] = None
    order_id: Optional[str] = None


class TradeControlHandler:
    """Approval state and execution readiness for semi-automated trading."""

    ORDER_TYPES = {
        "STOPLOSS_LIMIT_BUY",
        "MARKET_BUY",
        "LIMIT_BUY",
        "BRACKET_ORDER",
        "COVER_ORDER",
    }
    VALIDITY_TYPES = {"INTRADAY", "CARRY_FORWARD"}
    MODES = {"REGULAR", "BRACKET", "COVER"}
    _file_lock = threading.Lock()

    def __init__(self, backup_dir: str = "./trade_logs") -> None:
        self._lock = threading.Lock()
        self._counter = 0
        self._session_prefix = datetime.now(UTC).strftime("%Y%m%d%H%M%S")
        self._requests: Dict[str, TradeRequest] = {}
        self.backup_dir = Path(backup_dir)
        self.backup_dir.mkdir(parents=True, exist_ok=True)

    def create_trade_request(self, signal: dict, option_data: dict, quantity: int = 1) -> TradeRequest:
        with self._lock:
            self._counter += 1
            trade_id = f"TRADE_{self._session_prefix}_{self._counter:03d}"
            request = TradeRequest(
                trade_id=trade_id,
                symbol=signal["symbol"],
                signal=signal["signal"],
                entry=float(signal["entry"]),
                stop_loss=float(signal["stop_loss"]),
                targets=[float(target) for target in signal.get("targets", [])],
                category=signal.get("category", "index_options"),
                option_symbol=str(option_data.get("option_symbol", signal["symbol"])),
                option_type=str(option_data.get("option_type", signal.get("option_type", ""))),
                quantity=max(1, int(quantity)),
            )
            self._requests[trade_id] = request
            self._log_trade_event(request, "REQUESTED")
            return request

    def get(self, trade_id: str) -> Optional[TradeRequest]:
        return self._requests.get(trade_id)

    def set_order_preferences(
        self,
        trade_id: str,
        order_type: str | None = None,
        validity: str | None = None,
        mode: str | None = None,
    ) -> Optional[TradeRequest]:
        with self._lock:
            request = self._requests.get(trade_id)
            if not request or request.status != "PENDING":
                return request

            if order_type and order_type in self.ORDER_TYPES:
                request.order_type = order_type
            if validity and validity in self.VALIDITY_TYPES:
                request.validity = validity
            if mode and mode in self.MODES:
                request.mode = mode

            self._log_trade_event(request, "UPDATED")
            return request

    def approve_trade(self, trade_id: str) -> Optional[TradeRequest]:
        with self._lock:
            request = self._requests.get(trade_id)
            if not request or request.status != "PENDING":
                return request
            request.status = "APPROVED"
            request.approved_at = datetime.now(UTC).isoformat()
            self._log_trade_event(request, "APPROVED")
            return request

    def reject_trade(self, trade_id: str, reason: str = "") -> Optional[TradeRequest]:
        with self._lock:
            request = self._requests.get(trade_id)
            if not request or request.status != "PENDING":
                return request
            request.status = "REJECTED"
            request.rejection_reason = reason
            self._log_trade_event(request, "REJECTED")
            return request

    def cancel_trade(self, trade_id: str) -> Optional[TradeRequest]:
        with self._lock:
            request = self._requests.get(trade_id)
            if not request or request.status not in {"PENDING", "APPROVED"}:
                return request
            request.status = "CANCELLED"
            self._log_trade_event(request, "CANCELLED")
            return request

    def mark_executed(self, trade_id: str, order_id: str) -> Optional[TradeRequest]:
        with self._lock:
            request = self._requests.get(trade_id)
            if not request or request.status != "APPROVED":
                return request
            request.status = "EXECUTED"
            request.order_id = order_id
            request.executed_at = datetime.now(UTC).isoformat()
            self._log_trade_event(request, "EXECUTED")
            return request

    def pending_requests(self) -> List[TradeRequest]:
        with self._lock:
            return [request for request in self._requests.values() if request.status == "PENDING"]

    def approved_for_execution(self) -> List[TradeRequest]:
        with self._lock:
            return [request for request in self._requests.values() if request.status == "APPROVED"]

    def summary(self) -> dict:
        with self._lock:
            statuses = {"PENDING": 0, "APPROVED": 0, "REJECTED": 0, "CANCELLED": 0, "EXECUTED": 0}
            for request in self._requests.values():
                statuses[request.status] = statuses.get(request.status, 0) + 1
            return statuses

    def _log_trade_event(self, request: TradeRequest, event: str) -> None:
        with self._file_lock:
            file_path = self.backup_dir / f"trades_{datetime.now(UTC).strftime('%Y%m%d')}.xlsx"
            workbook = load_workbook(file_path) if file_path.exists() else Workbook()
            sheet = workbook.active
            sheet.title = "trade_log"

            if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
                sheet.append(
                    [
                        "timestamp",
                        "event",
                        "trade_id",
                        "symbol",
                        "signal",
                        "entry",
                        "stop_loss",
                        "target_1",
                        "target_2",
                        "target_3",
                        "quantity",
                        "order_type",
                        "validity",
                        "mode",
                        "status",
                        "order_id",
                    ]
                )

            targets = (request.targets + [None, None, None])[:3]
            sheet.append(
                [
                    datetime.now(UTC).isoformat(),
                    event,
                    request.trade_id,
                    request.symbol,
                    request.signal,
                    request.entry,
                    request.stop_loss,
                    targets[0],
                    targets[1],
                    targets[2],
                    request.quantity,
                    request.order_type,
                    request.validity,
                    request.mode,
                    request.status,
                    request.order_id,
                ]
            )
            workbook.save(file_path)
