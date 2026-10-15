"""Private Excel reporting rebuilt from the durable signal and live-fill databases."""

from __future__ import annotations

import json
import hashlib
import logging
import math
import os
import sqlite3
import tempfile
import threading
import time
import zipfile
from collections import defaultdict
from contextlib import contextmanager
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel

logger = logging.getLogger(__name__)
IST = timezone(timedelta(hours=5, minutes=30), name="IST")
SHEET_CATEGORIES = {
    "Index": "index_options",
    "Stocks": "nifty50_stock_options",
    "Commodity": "commodity_options",
}
SHEET_NAMES = ("Index", "Stocks", "Commodity", "Live Trades")
MAX_ROWS = 20000
MAX_EXPORT_BYTES = 50 * 1024 * 1024
MAX_TEXT_LENGTH = 2048
HEADERS = (
    "Serial No", "Trade Date", "Name", "Strike", "Expiry", "CE/PE",
    "Contract Symbol", "Entry Time (IST)", "Exit Time (IST)", "Total Duration",
    "Entry Price", "Observed Trigger Price", "Exit Price", "P&L (Points)",
    "Gross P&L (INR)", "Net P&L (INR)", "Qty Lots", "Lot Size", "Units",
    "Win % (running gross)", "Target Hit", "Stop Hit", "Session Auto-Square-Off",
    "Amount Used (INR)", "Amount After Exit (INR)", "Charges (INR)",
    "Charges Provenance/Status", "Entry Order Type", "Exit Order Type",
    "Status", "Accounting Basis", "Exit Reason", "Time Zone", "Notes", "Review",
    "Row ID", "Bot Data SHA-256",
)
HEADER_ROW = 6
ROW_ID_COLUMN = HEADERS.index("Row ID") + 1
REVIEW_COLUMNS = {"Notes": HEADERS.index("Notes") + 1, "Review": HEADERS.index("Review") + 1}
GREEN = PatternFill("solid", fgColor="E2F0D9")
RED = PatternFill("solid", fgColor="FCE4D6")
BLUE = PatternFill("solid", fgColor="1F4E78")


def _positive_int(value):
    if isinstance(value, bool):
        return None
    try:
        result = int(value)
    except (TypeError, ValueError, OverflowError):
        return None
    return result if result > 0 and result == value else None


def _positive_number(value):
    if isinstance(value, bool):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    return result if math.isfinite(result) and result > 0 else None


def _finite_number(value):
    if isinstance(value, bool):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    return result if math.isfinite(result) else None


def _epoch(value):
    try:
        result = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    if not math.isfinite(result) or result <= 0:
        return None
    return result / 1000 if result > 1e11 else result


def _local_datetime(value):
    stamp = _epoch(value)
    return datetime.fromtimestamp(stamp, IST).replace(tzinfo=None) if stamp else None


def _text(value):
    if value is None:
        return None
    text = str(value)
    if len(text) > MAX_TEXT_LENGTH:
        raise ValueError("journal text cell exceeds the safe size limit")
    if text[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + text
    return text


def _checksum_value(value):
    if isinstance(value, (datetime, date)):
        return {"excel_datetime": round(to_excel(value), 10)}
    if isinstance(value, timedelta):
        return {"excel_duration": round(value.total_seconds() / 86400, 10)}
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return format(float(value), ".15g")
    if isinstance(value, str):
        return _text(value)
    return value


def bot_row_checksum(values):
    """Hash bot-owned row values so the sync helper can reject edited source cells."""
    ignored = {HEADERS.index("Notes"), HEADERS.index("Review"), HEADERS.index("Bot Data SHA-256")}
    material = [_checksum_value(value) for index, value in enumerate(values) if index not in ignored]
    encoded = json.dumps(material, ensure_ascii=False, allow_nan=False, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _as_amount(price, quantity, multiplier):
    if price is None or quantity is None or multiplier is None:
        return None
    result = price * quantity * multiplier
    return round(result, 2) if math.isfinite(result) else None


def _read_only(path):
    path = Path(path)
    if not path.is_file():
        return None
    uri = path.resolve().as_uri() + "?mode=ro"
    connection = sqlite3.connect(uri, uri=True, timeout=5)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA query_only=ON")
    return connection


def _records(connection, table, limit=MAX_ROWS + 1):
    if connection is None:
        return []
    exists = connection.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,)
    ).fetchone()
    if not exists:
        return []
    rows = connection.execute(f"SELECT data FROM {table} LIMIT ?", (limit,)).fetchall()
    if len(rows) >= limit:
        raise ValueError("journal row limit exceeded")
    return [json.loads(row["data"]) for row in rows]


def _signal_record(position):
    meta = position.get("metadata") if isinstance(position.get("metadata"), dict) else {}
    lot_size = _positive_int(position.get("lot_size"))
    multiplier = _positive_number(meta.get("premium_multiplier"))
    entry = _finite_number(position.get("entry"))
    observed = _finite_number(position.get("observed_entry"))
    closed = position.get("status") in ("TARGET", "STOP", "SESSION_CLOSE")
    exit_price = _finite_number(position.get("exit_premium")) if closed else None
    exit_time = _local_datetime(position.get("closed_at")) if closed else None
    entry_time = _local_datetime(position.get("opened_at"))
    duration = (exit_time - entry_time) if exit_time and entry_time else None
    points = round(exit_price - entry, 4) if exit_price is not None and entry is not None else None
    units = lot_size
    gross = _as_amount(points, units, multiplier)
    entry_amount = _as_amount(entry, units, multiplier)
    exit_amount = _as_amount(exit_price, units, multiplier) if closed else None
    basis = "Hypothetical signal; one validated metadata lot; no order submitted"
    if multiplier is None:
        basis += "; INR conversion unavailable (premium_multiplier missing)"
    return {
        "row_id": "signal:" + str(position["id"]),
        "date": entry_time.date() if entry_time else None,
        "name": position.get("underlying"),
        "strike": _finite_number(position.get("strike")),
        "expiry": position.get("expiry"),
        "option": position.get("option_type"),
        "contract": position.get("option_symbol"),
        "entry_time": entry_time,
        "exit_time": exit_time,
        "duration": duration,
        "entry": entry,
        "observed": observed,
        "exit": exit_price,
        "points": points,
        "gross": gross,
        "net": None,
        "lots": 1 if lot_size else None,
        "lot_size": lot_size,
        "units": units,
        "target": "Yes" if position.get("status") == "TARGET" else "No",
        "stop": "Yes" if position.get("status") == "STOP" else "No",
        "session": "Yes" if position.get("status") == "SESSION_CLOSE" else "No",
        "amount_used": entry_amount,
        "amount_after": exit_amount,
        "charges": None,
        "charges_status": "N/A — hypothetical; no broker order",
        "entry_order": "N/A — no order submitted",
        "exit_order": "N/A — no order submitted" if closed else None,
        "status": position.get("status", "UNKNOWN"),
        "basis": basis,
        "exit_reason": position.get("status") if closed else None,
        "entry_epoch": _epoch(position.get("opened_at")),
        "exit_epoch": _epoch(position.get("closed_at")) if closed else None,
        "serial": None,
        "win_pct": None,
        "notes": None,
        "review": None,
    }


def _live_records(items):
    buys = {item.get("id"): item for item in items if item.get("side") == "BUY"}
    exits_by_entry = defaultdict(list)
    for item in items:
        if item.get("side") == "SELL" and item.get("entry_id") in buys:
            exits_by_entry[item["entry_id"]].append(item)

    for item in buys.values():
        if item.get("entry_notified") is not True or _positive_int(item.get("filled_quantity")) is None:
            continue
        entry_qty = _positive_int(item.get("filled_quantity"))
        entry_price = _positive_number(item.get("average_price"))
        entry_epoch = _epoch(item.get("first_fill_timestamp"))
        lot_size = _positive_int(item.get("lot_size"))
        fills = []
        own_exit_qty = _positive_int(item.get("exit_filled_quantity")) or 0
        own_exit_price = _positive_number(item.get("exit_average_price"))
        own_exit_time = _epoch(item.get("exit_fill_timestamp"))
        if own_exit_qty and own_exit_price:
            fills.append((own_exit_qty, own_exit_price, own_exit_time, None))
        for child in exits_by_entry.get(item.get("id"), []):
            quantity = _positive_int(child.get("filled_quantity"))
            price = _positive_number(child.get("average_price"))
            if quantity and price:
                fills.append((
                    quantity, price, _epoch(child.get("fill_timestamp")),
                    str(child.get("order_type") or ""),
                ))
        exit_qty = sum(fill[0] for fill in fills)
        if exit_qty > entry_qty:
            raise ValueError("broker exit quantity exceeds confirmed entry quantity")
        closed = (
            exit_qty == entry_qty and entry_qty > 0 and entry_price is not None
            and entry_epoch is not None and bool(fills)
            and all(fill[2] is not None for fill in fills)
        )
        exit_price = (
            sum(quantity * price for quantity, price, _, _ in fills) / exit_qty
            if closed and exit_qty else None
        )
        exit_epochs = [fill[2] for fill in fills if fill[2] is not None]
        exit_epoch = max(exit_epochs) if closed and exit_epochs else None
        entry_time = _local_datetime(entry_epoch)
        exit_time = _local_datetime(exit_epoch)
        duration = exit_time - entry_time if exit_time and entry_time else None
        points = round(exit_price - entry_price, 4) if closed and exit_price is not None else None
        gross = round(points * entry_qty, 2) if points is not None else None
        triggered = set(item.get("legs_triggered") or [])
        exit_reason = item.get("exit_reason")
        if not exit_reason and triggered:
            exit_reason = (
                "time_exit_1525" if "TARGET_LEG" in triggered and item.get("time_exit")
                else "target" if "TARGET_LEG" in triggered
                else "stop_loss" if "STOP_LOSS_LEG" in triggered
                else "broker_exit_unattributed"
            )
        if not exit_reason:
            exit_reason = next((child.get("reason") for child in exits_by_entry.get(item.get("id"), [])
                                if child.get("reason")), None)
        exit_order = None
        if closed:
            if "STOP_LOSS_LEG" in triggered:
                exit_order = "MARKET"
            elif "TARGET_LEG" in triggered:
                exit_order = "LIMIT"
            else:
                types = {fill[3].upper() for fill in fills if fill[3]}
                exit_order = next(iter(types)) if len(types) == 1 and next(iter(types)) in {
                    "LIMIT", "MARKET",
                } else None
        expiry = item.get("expiry")
        yield {
            "row_id": "live:" + str(item["id"]),
            "date": entry_time.date() if entry_time else None,
            "name": item.get("underlying") or "NIFTY",
            "strike": _finite_number(item.get("strike")),
            "expiry": expiry,
            "option": item.get("option_type"),
            "contract": item.get("option_symbol") or item.get("symbol"),
            "entry_time": entry_time,
            "exit_time": exit_time,
            "duration": duration,
            "entry": entry_price,
            "observed": _finite_number(item.get("breakout_price")),
            "exit": exit_price,
            "points": points,
            "gross": gross,
            "net": None,
            "lots": entry_qty // lot_size if lot_size and entry_qty % lot_size == 0 else None,
            "lot_size": lot_size,
            "units": entry_qty,
            "target": "Yes" if (
                closed and "TARGET_LEG" in triggered
                and exit_reason != "time_exit_1525" and not item.get("time_exit")
            ) else "No",
            "stop": "Yes" if closed and "STOP_LOSS_LEG" in triggered else "No",
            "session": "Yes" if closed and exit_reason == "time_exit_1525" else "No",
            "amount_used": round(entry_price * entry_qty, 2) if entry_price is not None else None,
            "amount_after": round(exit_price * exit_qty, 2) if closed and exit_price is not None else None,
            "charges": None,
            "charges_status": "Pending contract note/ledger; actual charges unavailable",
            "entry_order": (
                item.get("entry_order_type") or item.get("order_type")
            ) if (item.get("entry_order_type") or item.get("order_type")) in {
                "LIMIT", "MARKET",
            } else None,
            "exit_order": exit_order,
            "status": "CLOSED" if closed else ("OPEN" if item.get("owned_quantity", 0) else "UNKNOWN"),
            "basis": "Broker-confirmed fills; Dhan average premium × confirmed quantity; actual trade only",
            "exit_reason": exit_reason if closed else None,
            "entry_epoch": entry_epoch,
            "exit_epoch": exit_epoch,
            "serial": None,
            "win_pct": None,
            "notes": None,
            "review": None,
        }


def _classify(records):
    wins = losses = breakeven = unknown = 0
    daily = defaultdict(lambda: {
        "wins": 0, "losses": 0, "breakeven": 0, "unknown": 0, "gross": 0.0, "priced": 0,
    })
    for record in records:
        if record["status"] not in ("TARGET", "STOP", "SESSION_CLOSE", "CLOSED"):
            continue
        close_day = record["exit_time"].date() if record["exit_time"] else record["date"]
        day = close_day.isoformat() if close_day else "Unknown"
        if record["gross"] is None:
            unknown += 1
            daily[day]["unknown"] += 1
        else:
            daily[day]["gross"] += record["gross"]
            daily[day]["priced"] += 1
            if record["gross"] > 0:
                wins += 1
                daily[day]["wins"] += 1
            elif record["gross"] < 0:
                losses += 1
                daily[day]["losses"] += 1
            else:
                breakeven += 1
                daily[day]["breakeven"] += 1
    denominator = wins + losses + breakeven
    win_rate = wins / denominator if denominator else None
    return wins, losses, breakeven, unknown, win_rate, daily


def _cell(sheet, value, number_format=None, fill=None, bold=False):
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.value = _text(value)
        cell.number_format = "@"
    if number_format:
        cell.number_format = number_format
    if fill:
        cell.fill = fill
    if bold:
        cell.font = Font(bold=True)
    return cell


def _output_values(record, serial, win_pct):
    values = [
        serial, record.get("date"), record.get("name"), record.get("strike"),
        record.get("expiry"), record.get("option"), record.get("contract"),
        record.get("entry_time"), record.get("exit_time"), record.get("duration"),
        record.get("entry"), record.get("observed"), record.get("exit"),
        record.get("points"), record.get("gross"), record.get("net"),
        record.get("lots"), record.get("lot_size"), record.get("units"),
        win_pct, record.get("target"), record.get("stop"), record.get("session"),
        record.get("amount_used"), record.get("amount_after"), record.get("charges"),
        record.get("charges_status"), record.get("entry_order"), record.get("exit_order"),
        record.get("status"), record.get("basis"), record.get("exit_reason"), "IST",
        record.get("notes"), record.get("review"), record.get("row_id"),
    ]
    values.append(bot_row_checksum(values))
    return values


class TradeJournalExporter:
    """Regenerate one four-sheet workbook from authoritative persisted state."""

    def __init__(self, signal_db_path, live_db_path, output_path, *, clock=time.time):
        self.signal_db_path = str(signal_db_path)
        self.live_db_path = str(live_db_path)
        self.output_path = Path(output_path).expanduser()
        if self.output_path.suffix.lower() != ".xlsx":
            raise ValueError("trade journal output must use the .xlsx extension")
        self.clock = clock
        self._export_lock = threading.Lock()
        self.last_export_at = None
        self.last_error = None
        self.last_rows = 0
        self.last_event_at = None
        self.pending_rows = 0

    def _snapshot(self):
        signal_db = _read_only(self.signal_db_path)
        live_db = _read_only(self.live_db_path)
        try:
            if signal_db is None:
                raise FileNotFoundError("signal lifecycle database unavailable")
            if not signal_db.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='lifecycle_signals'"
            ).fetchone():
                raise ValueError("signal lifecycle table unavailable")
            signals = _records(signal_db, "lifecycle_signals")
            intents = _records(live_db, "live_intents")
        finally:
            if signal_db is not None:
                signal_db.close()
            if live_db is not None:
                live_db.close()

        signal_symbols = {
            str(position["id"]): position.get("option_symbol")
            for position in signals
            if position.get("id") is not None and isinstance(position.get("option_symbol"), str)
        }
        for item in intents:
            if not item.get("option_symbol"):
                symbol = signal_symbols.get(str(item.get("signal_key")))
                if symbol:
                    item["option_symbol"] = symbol

        categories = {name: [] for name in SHEET_CATEGORIES}
        for position in signals:
            category = position.get("category")
            sheet = next((name for name, value in SHEET_CATEGORIES.items() if value == category), None)
            if sheet is None:
                continue
            categories[sheet].append(_signal_record(position))
        for records in categories.values():
            records.sort(key=lambda row: (row["entry_epoch"] or 0, row["row_id"]))
        live = list(_live_records(intents))
        live.sort(key=lambda row: (row["entry_epoch"] or 0, row["row_id"]))
        total = sum(map(len, categories.values())) + len(live)
        if total > MAX_ROWS:
            raise ValueError("journal row limit exceeded")
        all_records = [record for records in categories.values() for record in records] + live
        self.last_event_at = max(
            (record["exit_epoch"] or record["entry_epoch"] or 0 for record in all_records),
            default=None,
        )
        self.pending_rows = sum(
            record["status"] in ("OPEN", "UNKNOWN") for record in all_records
        )
        return categories, live

    def _write_sheet(self, sheet, name, records, generated):
        wins, losses, breakeven, unknown, win_rate, daily = _classify(records)
        net_known = sum(record["net"] is not None for record in records)
        sheet.append([_cell(sheet, f"Trade Journal — {name}", bold=True)])
        sheet.append([
            _cell(sheet, "Report generated (IST)", bold=True),
            _cell(sheet, generated, "dd-mmm-yyyy hh:mm:ss"),
        ])
        basis = (
            "Hypothetical signal outcomes; no broker execution. INR amounts require "
            "explicit premium_multiplier metadata. Live Trades are separate broker-confirmed fills."
            if name != "Live Trades" else
            "Actual broker-confirmed fills only. Gross P&L from confirmed prices/quantities; "
            "net remains unavailable until verified charges are sourced."
        )
        sheet.append([_cell(sheet, "Accounting basis", bold=True), _cell(sheet, basis)])
        rate_text = "N/A" if win_rate is None else f"{win_rate:.1%}"
        sheet.append([_cell(
            sheet,
            f"Closed priced outcomes — wins {wins}; losses {losses}; break-even {breakeven}; "
            f"unpriced {unknown}; gross win rate {rate_text}; net priced {net_known}. "
            "Open/unknown trades excluded. Break-even remains in the denominator.",
        )])
        closed_records = [
            record for record in records
            if record["status"] in ("TARGET", "STOP", "SESSION_CLOSE", "CLOSED")
        ]
        priced_gross = [record["gross"] for record in closed_records if record["gross"] is not None]
        gross_total = f"₹{sum(priced_gross):,.2f}" if priced_gross else "N/A (no priced closed results)"
        gross_profit = sum(value for value in priced_gross if value > 0)
        gross_loss = sum(value for value in priced_gross if value < 0)
        gross_profit_text = f"₹{gross_profit:,.2f}" if priced_gross else "N/A"
        gross_loss_text = f"₹{gross_loss:,.2f}" if priced_gross else "N/A"
        net_values = [record["net"] for record in closed_records]
        net_total = (
            f"₹{sum(net_values):,.2f}" if net_values and all(value is not None for value in net_values)
            else "N/A (verified charges unavailable)"
        )
        amount_used = [record["amount_used"] for record in records if record["amount_used"] is not None]
        amount_after = [record["amount_after"] for record in closed_records if record["amount_after"] is not None]
        amount_used_text = f"₹{sum(amount_used):,.2f}" if amount_used else "N/A"
        amount_after_text = f"₹{sum(amount_after):,.2f}" if amount_after else "N/A"
        sheet.append([_cell(
            sheet,
            f"Gross P&L (priced closed): {gross_total} | Gross profit: {gross_profit_text} | "
            f"Gross loss: {gross_loss_text} | Net P&L: {net_total} | "
            f"Premium amount used (known): {amount_used_text} | "
            f"Exit proceeds before charges (known): {amount_after_text}",
        )])
        sheet.append([_cell(sheet, header, fill=BLUE, bold=True) for header in HEADERS])
        sheet.freeze_panes = f"A{HEADER_ROW + 1}"
        sheet.sheet_view.showGridLines = False
        sheet.column_dimensions[get_column_letter(ROW_ID_COLUMN)].width = 34
        widths = {
            2: 14, 3: 20, 7: 32, 8: 21, 9: 21, 10: 16, 27: 42,
            30: 18, 31: 70, 32: 26, 34: 36, 35: 36,
        }
        for index, width in widths.items():
            sheet.column_dimensions[get_column_letter(index)].width = width

        running_wins = running_losses = running_be = 0
        running_rates = {}
        closed_priced = sorted(
            (
                record for record in records
                if record["status"] in ("TARGET", "STOP", "SESSION_CLOSE", "CLOSED")
                and record["gross"] is not None
            ),
            key=lambda record: (record["exit_epoch"] or 0, record["row_id"]),
        )
        for closed_record in closed_priced:
            if closed_record["gross"] > 0:
                running_wins += 1
            elif closed_record["gross"] < 0:
                running_losses += 1
            else:
                running_be += 1
            running_rates[closed_record["row_id"]] = (
                running_wins / (running_wins + running_losses + running_be)
            )
        for serial, record in enumerate(records, 1):
            row_rate = running_rates.get(record["row_id"])
            record["serial"] = serial
            record["win_pct"] = row_rate
            values = _output_values(record, serial, row_rate)
            row = []
            for col, value in enumerate(values, 1):
                number_format = None
                fill = None
                if col in (2,):
                    number_format = "dd-mmm-yyyy"
                elif col in (8, 9):
                    number_format = "dd-mmm-yyyy hh:mm:ss"
                elif col == 10:
                    number_format = "[h]:mm:ss"
                elif col in (11, 12, 13, 14, 15, 16, 24, 25, 26):
                    number_format = '#,##0.00;[Red](#,##0.00);-'
                elif col == 20:
                    number_format = "0.0%"
                elif col == ROW_ID_COLUMN:
                    number_format = "@"
                if col in (15, 16) and isinstance(value, (int, float)):
                    fill = GREEN if value > 0 else RED if value < 0 else None
                row.append(_cell(sheet, value, number_format, fill))
            sheet.append(row)
        end_row = HEADER_ROW + len(records)
        sheet.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(HEADERS))}{max(end_row, HEADER_ROW)}"
        sheet.append([_cell(sheet, "Daily closed results — gross, priced outcomes only", bold=True)])
        sheet.append([
            _cell(sheet, value, fill=BLUE, bold=True)
            for value in (
                "Date", "Wins", "Losses", "Break-even", "Unpriced closed",
                "Gross win rate", "Gross P&L (INR)",
            )
        ])
        for day, counts in sorted(daily.items()):
            denominator = counts["wins"] + counts["losses"] + counts["breakeven"]
            rate = counts["wins"] / denominator if denominator else None
            sheet.append([
                _cell(sheet, day, "@"), _cell(sheet, counts["wins"]),
                _cell(sheet, counts["losses"]), _cell(sheet, counts["breakeven"]),
                _cell(sheet, counts["unknown"]), _cell(sheet, rate, "0.0%"),
                _cell(sheet, round(counts["gross"], 2) if counts["priced"] else None, '#,##0.00;[Red](#,##0.00);-'),
            ])
        return len(records)

    def export(self):
        if not self._export_lock.acquire(blocking=False):
            return False
        try:
            categories, live = self._snapshot()
            generated = datetime.fromtimestamp(self.clock(), IST).replace(tzinfo=None)
            workbook = Workbook(write_only=True)
            for name in SHEET_NAMES:
                sheet = workbook.create_sheet(name)
                records = live if name == "Live Trades" else categories[name]
                self._write_sheet(sheet, name, records, generated)

            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            with _file_lock(self.output_path.with_suffix(self.output_path.suffix + ".lock")):
                fd, temp_name = tempfile.mkstemp(
                    prefix=".trade-journal-", suffix=".xlsx", dir=str(self.output_path.parent)
                )
                os.close(fd)
                temp_path = Path(temp_name)
                try:
                    os.chmod(temp_path, 0o600)
                    workbook.save(temp_path)
                    if temp_path.stat().st_size > MAX_EXPORT_BYTES or not zipfile.is_zipfile(temp_path):
                        raise ValueError("generated workbook failed size/integrity validation")
                    with zipfile.ZipFile(temp_path) as archive:
                        if archive.testzip() is not None:
                            raise ValueError("generated workbook failed integrity validation")
                    self._replace_with_retry(temp_path, self.output_path)
                finally:
                    temp_path.unlink(missing_ok=True)
            self.last_export_at = self.clock()
            self.last_rows = sum(map(len, categories.values())) + len(live)
            self.last_error = None
            return True
        except Exception as exc:
            self.last_error = type(exc).__name__
            logger.warning("Trade journal export failed: %s", type(exc).__name__)
            raise
        finally:
            self._export_lock.release()

    @staticmethod
    def _replace_with_retry(source, target):
        for attempt in range(5):
            try:
                os.replace(source, target)
                return
            except PermissionError:
                if attempt == 4:
                    raise
                time.sleep(0.2 * (attempt + 1))

    def status(self):
        return {
            "last_event_at": self.last_event_at,
            "last_export_at": self.last_export_at,
            "pending_rows": self.pending_rows,
            "exported_rows": self.last_rows,
            "export_error": self.last_error,
        }


@contextmanager
def _file_lock(path, timeout=10):
    """Cross-platform advisory lock shared by the server exporter and local sync."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    handle = open(path, "a+b")
    try:
        handle.seek(0, os.SEEK_END)
        if handle.tell() == 0:
            handle.write(b"\0")
            handle.flush()
        deadline = time.monotonic() + timeout
        while True:
            try:
                if os.name == "nt":
                    import msvcrt

                    handle.seek(0)
                    msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    import fcntl

                    fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                break
            except (BlockingIOError, OSError):
                if time.monotonic() >= deadline:
                    raise TimeoutError("trade journal file is busy")
                time.sleep(0.1)
        yield
    finally:
        try:
            if os.name == "nt":
                import msvcrt

                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl

                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
        except OSError:
            pass
        handle.close()


class TradeJournalWorker:
    """Coalesce database changes and rebuild outside scanner/order critical paths."""

    def __init__(self, exporter, interval=5):
        self.exporter = exporter
        self.interval = max(1, int(interval))
        self._stop = threading.Event()
        self._wake = threading.Event()
        self._thread = None
        self._fingerprint = None
        self.export_pending = True

    def _source_fingerprint(self):
        fingerprint = []
        for path in (self.exporter.signal_db_path, self.exporter.live_db_path):
            for candidate in (Path(path), Path(str(path) + "-wal")):
                try:
                    stat = candidate.stat()
                    fingerprint.append((str(candidate), stat.st_mtime_ns, stat.st_size))
                except OSError:
                    fingerprint.append((str(candidate), None, None))
        return tuple(fingerprint)

    def run_once(self, force=False):
        current = self._source_fingerprint()
        if force or self.export_pending or current != self._fingerprint:
            self.export_pending = True
            if not self.exporter.export():
                return False
            self._fingerprint = current
            self.export_pending = False
            return True
        return False

    def _run(self):
        self._wake.set()
        while not self._stop.is_set():
            self._wake.wait(self.interval)
            self._wake.clear()
            if self._stop.is_set():
                break
            try:
                self.run_once()
            except Exception:
                pass

    def start(self):
        if self._thread and self._thread.is_alive():
            return
        self._stop.clear()
        self._thread = threading.Thread(target=self._run, name="trade-journal-export", daemon=True)
        self._thread.start()

    def wake(self):
        self._wake.set()

    def stop(self):
        self._stop.set()
        self._wake.set()
        if self._thread and self._thread.is_alive():
            self._thread.join(timeout=2)

    def status(self):
        result = self.exporter.status()
        result["running"] = bool(self._thread and self._thread.is_alive())
        result["export_pending"] = self.export_pending
        return result
